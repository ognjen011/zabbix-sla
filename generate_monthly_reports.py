#!/usr/bin/env python3
"""Generate separate complete-month SLA/usage reports plus a trend comparison."""
import argparse
from datetime import datetime
from pathlib import Path

from openpyxl.chart import LineChart, Reference

from report_retention import retain_monthly_workbook

from link_usage import calendar_months, collect_link_report, add_usage_sheet, validate_config
from zabbix_sla_report import ZabbixAPI, ExcelReportGenerator, load_config


def add_availability_sheet(report, group_name, rows):
    sheet = report.workbook.create_sheet(f"SLA {group_name}"[:31], 0)
    headers = ["Host", "Technical Host", "Month", "Availability (%)", "SLA Target (%)", "Status", "Downtime (seconds)"]
    sheet.append(headers)
    for row in rows:
        sheet.append([row[key] for key in ('name', 'host', 'month', 'device_sla', 'sla_target', 'sla_status', 'downtime_seconds')])
        fill, font = report.get_cell_style(row['device_sla'])
        for index in (4, 6):
            sheet.cell(sheet.max_row, index).fill = fill
            sheet.cell(sheet.max_row, index).font = font
        sheet.cell(sheet.max_row, 4).number_format = '0.00'
    for cell in sheet[1]:
        cell.fill, cell.font = report.header_fill, report.header_font
        sheet.column_dimensions[cell.column_letter].width = 25
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def generate_reports(api, config, months, output_dir):
    settings = dict(config.get('link_usage', {}) or {})
    settings['history_months'] = months
    config = {**config, 'link_usage': settings}
    validate_config(config)
    periods = calendar_months(int(datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0).timestamp()) - 1, months)
    groups_config = config.get('host_groups', {}) or {}
    groups = api.get_host_groups(list(groups_config))
    if not groups:
        raise ValueError('No configured groups found in Zabbix')
    missing = set(groups_config) - {group['name'] for group in groups}
    if missing:
        raise ValueError(f'Configured groups missing from Zabbix: {sorted(missing)}')
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    files = []
    for group in groups:
        name = group['name']
        group_config = groups_config[name] or {}
        target = group_config.get('sla_threshold', config.get('default_sla_threshold', 99.9))
        warning = group_config.get('orange_threshold', config.get('default_orange_threshold', 5))
        excluded = {host.lower() for host in (config.get('global_excluded_hosts') or []) + (group_config.get('excluded_hosts') or [])}
        hosts = [host for host in api.get_hosts_in_group(group['groupid']) if host['name'].lower() not in excluded and host['host'].lower() not in excluded]
        if not hosts:
            raise ValueError(f'No enabled, included hosts in group {name}')
        print(f'{name}: {len(hosts)} hosts; {months} complete months', flush=True)
        all_hosts = []
        for index, host in enumerate(hosts, 1):
            print(f'  [{index}/{len(hosts)}] {host["name"]}', flush=True)
            links = collect_link_report(api, host['hostid'], periods[-1][0], periods[-1][1] - 1, config)
            availability = []
            for start, stop in periods:
                stats = api.get_host_availability(host['hostid'], start, stop - 1)
                sla = stats['availability']
                availability.append({'name':host['name'], 'host':host['host'],
                    'month':datetime.fromtimestamp(start).strftime('%Y-%m'), 'device_sla':sla,
                    'sla_target':target, 'sla_status':'COMPLIANT' if sla >= target else 'WARNING' if sla >= target - warning else 'BREACH',
                    'downtime_seconds':stats['downtime_seconds'], 'total_seconds':stats['total_seconds']})
            all_hosts.append({**host, 'link_usage':links['latest'], 'link_history':links['monthly'], 'link_daily':links['daily'], 'availability':availability})
        safe_name = ''.join(character if character.isalnum() or character in '-_' else '_' for character in name)
        for month_index, (start, stop) in enumerate(periods):
            month = datetime.fromtimestamp(start).strftime('%Y-%m')
            report = ExcelReportGenerator(target, warning)
            sla_rows = [host['availability'][month_index] for host in all_hosts]
            sheet = add_availability_sheet(report, name, sla_rows)
            seconds = sum(row['total_seconds'] for row in sla_rows)
            downtime = sum(row['downtime_seconds'] for row in sla_rows)
            overall = (1 - downtime / seconds) * 100 if seconds else 100
            sheet.append(['Overall group', '', month, overall, target])
            monthly_hosts = [{**host, 'link_usage':[r for r in host['link_history'] if r['Month'] == month], 'link_history':[], 'link_daily':[r for r in host['link_daily'] if r['Date'].startswith(month)]} for host in all_hosts]
            add_usage_sheet(report, name, monthly_hosts)
            path = output_dir / f'SLA_Usage_{safe_name}_{month}.xlsx'
            report.workbook.save(path)
            files.append(path)
            if config.get("auto_save_reports", True):
                retain_monthly_workbook(path, name)
            print(f'Saved {path}; overall SLA {overall:.2f}%', flush=True)
        comparison = ExcelReportGenerator(target, warning)
        sla_rows = [row for host in all_hosts for row in host['availability']]
        add_availability_sheet(comparison, name, sla_rows)
        add_usage_sheet(comparison, name, all_hosts)
        summary = comparison.workbook.create_sheet('SLA Trends', 0)
        summary.append(['Month', 'Overall SLA (%)', 'SLA Target (%)', 'Compliant', 'Warning', 'Breach'])
        for month_index, (start, stop) in enumerate(periods):
            rows = [host['availability'][month_index] for host in all_hosts]
            seconds = sum(row['total_seconds'] for row in rows)
            downtime = sum(row['downtime_seconds'] for row in rows)
            overall = (1 - downtime / seconds) * 100 if seconds else 100
            summary.append([datetime.fromtimestamp(start).strftime('%Y-%m'), overall, target, *[sum(row['sla_status'] == status for row in rows) for status in ('COMPLIANT', 'WARNING', 'BREACH')]])
        chart = LineChart()
        chart.title = f'{name}: Monthly SLA'
        chart.y_axis.title = 'Availability (%)'
        chart.add_data(Reference(summary, min_col=2, max_col=3, min_row=1, max_row=months + 1), titles_from_data=True)
        chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=months + 1))
        summary.add_chart(chart, 'H2')
        path = output_dir / f'Trend_Comparison_{safe_name}_{datetime.fromtimestamp(periods[0][0]):%Y-%m}_to_{datetime.fromtimestamp(periods[-1][0]):%Y-%m}.xlsx'
        comparison.workbook.save(path)
        files.append(path)
        if config.get("auto_save_reports", True):
            retain_monthly_workbook(path, name)
        print(f'Saved trend comparison: {path}', flush=True)
    return files


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', default='config.yaml')
    parser.add_argument('--months', type=int, default=3)
    parser.add_argument('--output-dir', default='reports')
    args = parser.parse_args()
    config = load_config(args.config)
    api = ZabbixAPI(config['zabbix']['url'], config['zabbix']['token'])
    generate_reports(api, config, args.months, args.output_dir)


if __name__ == '__main__':
    main()
