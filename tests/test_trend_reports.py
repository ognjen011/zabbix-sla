import io
import unittest
from datetime import datetime, timedelta

from link_usage import calendar_months, collect_link_report, add_usage_sheet, validate_config
from tests.test_link_usage import FakeAPI


class HistoryAPI(FakeAPI):
    def __init__(self, missing_middle=False, missing_speed=False):
        super().__init__()
        self.calls = []
        self.missing_middle = missing_middle
        self.missing_speed = missing_speed

    def _call(self, method, params):
        self.calls.append((method, params))
        if method == 'item.get':
            return super()._call(method, params)
        trends = []
        for start, stop in calendar_months(params['time_till'], 3):
            month = datetime.fromtimestamp(start).month
            if self.missing_middle and month == 8:
                continue
            for clock in range(start, stop, 3600):
                for itemid in params['itemids']:
                    if itemid == '3' and self.missing_speed:
                        continue
                    value = month * (1e8 if itemid == '3' else 1e6)
                    trends.append({'itemid': itemid, 'clock': str(clock), 'num': '60', 'value_avg': str(value), 'value_max': str(value if itemid == '3' else value * 2)})
        return trends


def report(api):
    periods = calendar_months(int(datetime(2026, 9, 30, 23, 59, 59).timestamp()), 3)
    return collect_link_report(api, '1', periods[-1][0], periods[-1][1] - 1, {})


class TrendTests(unittest.TestCase):
    def test_calendar_months_cross_year_and_leap_day(self):
        periods = calendar_months(int(datetime(2024, 2, 29, 23, 59, 59).timestamp()), 3)
        self.assertEqual([datetime.fromtimestamp(start).strftime('%Y-%m') for start, _ in periods], ['2023-12', '2024-01', '2024-02'])
        self.assertEqual(datetime.fromtimestamp(periods[-1][1]), datetime(2024, 3, 1))

    def test_three_months_daily_peaks_and_historical_capacity(self):
        api = HistoryAPI()
        links = report(api)
        self.assertEqual(len(links['monthly']), 3)
        self.assertEqual(len(links['latest']), 1)
        self.assertEqual(len(links['daily']), 92)
        july, august, september = links['monthly']
        self.assertEqual(july['Download Avg (Mbps)'], 7)
        self.assertAlmostEqual(august['Download Avg Growth (%)'], 100 / 7)
        self.assertAlmostEqual(september['Capacity Growth (%)'], 12.5)
        self.assertEqual(september['Capacity (Mbps)'], 900)
        self.assertEqual(september['Download Avg (%)'], 1)
        self.assertEqual(links['daily'][0]['Download Peak (Mbps)'], 14)
        self.assertAlmostEqual(sum(row['Download (GB)'] for row in links['daily'] if row['Date'].startswith('2026-07')), july['Download (GB)'])
        self.assertEqual(sum(method == 'trend.get' for method, _ in api.calls), 1)

    def test_missing_middle_does_not_compare_non_adjacent_months(self):
        links = report(HistoryAPI(missing_middle=True))
        self.assertIsNone(links['monthly'][1]['Download (GB)'])
        self.assertIsNone(links['monthly'][2]['Download Avg Growth (%)'])
        self.assertIsNone(links['monthly'][2]['Capacity Growth (%)'])

    def test_no_historical_speed_does_not_invent_capacity_growth(self):
        links = report(HistoryAPI(missing_speed=True))
        self.assertEqual(links['monthly'][0]['Capacity Basis'], 'Current interface speed')
        self.assertEqual(links['monthly'][0]['Capacity (Mbps)'], 100)
        self.assertIsNone(links['monthly'][2]['Capacity Growth (%)'])

    def test_config_limits(self):
        for value in (0, 13, '3', True):
            with self.assertRaises(ValueError):
                validate_config({'link_usage': {'history_months': value}})

    def test_excel_contains_daily_monthly_and_charts(self):
        from zabbix_sla_report import ExcelReportGenerator
        from openpyxl import load_workbook
        links = report(HistoryAPI())
        generator = ExcelReportGenerator(99.9, 5)
        add_usage_sheet(generator, 'Example Group', [{'name': 'router', 'link_usage': links['latest'], 'link_history': links['monthly'], 'link_daily': links['daily']}])
        buffer = io.BytesIO()
        generator.workbook.save(buffer)
        buffer.seek(0)
        workbook = load_workbook(buffer)
        self.assertEqual(workbook.sheetnames, ['Links Example Group', 'Monthly Example Group', 'Daily Example Group'])
        self.assertEqual(workbook['Daily Example Group'].max_row, 93)
        self.assertEqual(len(workbook['Monthly Example Group']._charts), 3)


if __name__ == '__main__':
    unittest.main()
