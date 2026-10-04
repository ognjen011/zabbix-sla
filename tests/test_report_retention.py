import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

import database
from report_retention import retain_monthly_workbook


class RetentionTests(unittest.TestCase):
    def test_comparison_import_skips_sla_trends_sheet_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(database, 'DB_PATH', Path(directory) / 'test.db'):
            workbook = Workbook()
            summary = workbook.active
            summary.title = 'SLA Trends'
            summary.append(['Month', 'Overall SLA (%)'])
            summary.append(['2026-09', 99.8])
            sla = workbook.create_sheet('SLA Example Group')
            sla.append(['Host', 'Technical Host', 'Month', 'Availability (%)', 'SLA Target (%)', 'Status'])
            for month in ('2026-07', '2026-08', '2026-09'):
                sla.append(['router', 'technical-router', month, 99.8, 99.99, 'WARNING'])
            monthly = workbook.create_sheet('Monthly Example Group')
            monthly.append(['Host', 'Interface', 'Interface ID', 'Description', 'Month'])
            for month in ('2026-07', '2026-08', '2026-09'):
                monthly.append(['router', 'eth0', '7', 'PRI', month])
            path = Path(directory) / 'comparison.xlsx'
            workbook.save(path)
            first = retain_monthly_workbook(path, 'Example Group')
            second = retain_monthly_workbook(path, 'Example Group')
            self.assertEqual(first, second)
            self.assertEqual(database.get_report_count(), 1)
            report = database.get_report(first)
            self.assertEqual(report['summary_data'][0]['overall_sla'], 99.8)
            self.assertEqual(len(report['detail_data']['Example Group'][0]['link_history']), 3)


if __name__ == '__main__':
    unittest.main()
