import unittest
from link_usage import monthly_link_usage, validate_config, add_usage_sheet


class FakeAPI:
    def __init__(self, units="bps", description="PRI WAN"):
        self.units = units
        self.description = description

    def _call(self, method, params):
        if method == "item.get":
            return [
                {"itemid": "1", "key_": "net.if.in[ifHCInOctets.7]", "name": f"Interface eth7({self.description}): Bits received", "units": self.units},
                {"itemid": "2", "key_": "net.if.out[ifHCOutOctets.7]", "name": f"Interface eth7({self.description}): Bits sent", "units": self.units},
                {"itemid": "3", "key_": "net.if.speed[ifHighSpeed.7]", "name": "speed", "units": "bps", "lastvalue": "100000000"},
            ]
        return [{"itemid": "1", "clock": "0", "num": "60", "value_avg": "8000000", "value_max": "16000000"}] if "1" in params["itemids"] else []


class UsageTests(unittest.TestCase):
    def test_volume_peak_capacity_and_missing_direction(self):
        row = monthly_link_usage(FakeAPI(), "1", 0, 7199, {})[0]
        self.assertEqual(row["Download (GB)"], 3.6)
        self.assertEqual(row["Download Avg (Mbps)"], 8)
        self.assertEqual(row["Download Peak (Mbps)"], 16)
        self.assertEqual(row["Download Avg (%)"], 8)
        self.assertEqual(row["Download Trend Coverage (%)"], 50)
        self.assertIsNone(row["Upload (GB)"])

    def test_custom_description_regex(self):
        config = {"link_usage": {"description_regex": "uplink"}}
        self.assertEqual(monthly_link_usage(FakeAPI(), "1", 0, 3599, config), [])
        self.assertEqual(len(monthly_link_usage(FakeAPI(description="UPLINK"), "1", 0, 3599, config)), 1)
        self.assertEqual(monthly_link_usage(FakeAPI(description=""), "1", 0, 3599, {}), [])

    def test_byte_rates_and_unknown_units(self):
        row = monthly_link_usage(FakeAPI(units="Bps"), "1", 0, 3599, {})[0]
        self.assertEqual(row["Download (GB)"], 28.8)
        row = monthly_link_usage(FakeAPI(units="bytes"), "1", 0, 3599, {})[0]
        self.assertIsNone(row["Download (GB)"])

    def test_invalid_regex_and_disabled_collection(self):
        with self.assertRaises(ValueError):
            validate_config({"link_usage": {"description_regex": ""}})
        with self.assertRaises(ValueError):
            validate_config({"link_usage": {"description_regex": "["}})
        self.assertEqual(monthly_link_usage(FakeAPI(), "1", 0, 3599, {"link_usage": {"enabled": False}}), [])

    def test_local_token_file(self):
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from zabbix_sla_report import load_config, ZabbixAPI
        with TemporaryDirectory() as directory:
            path = Path(directory)
            (path / "secret").write_text(" test-token\n")
            (path / "config.yaml").write_text("zabbix:\n  token_file: secret\n")
            config = load_config(path / "config.yaml")
            self.assertEqual(config["zabbix"]["token"], "test-token")
            self.assertEqual(ZabbixAPI("http://localhost/zabbix", " token\n").token, "token")

    def test_excel(self):
        from zabbix_sla_report import ExcelReportGenerator
        from io import BytesIO
        from openpyxl import load_workbook
        report = ExcelReportGenerator(99.9, 5)
        rows = monthly_link_usage(FakeAPI(), "1", 0, 3599, {})
        add_usage_sheet(report, "WAN", [{"name": "router", "link_usage": rows}])
        buffer = BytesIO()
        report.workbook.save(buffer)
        buffer.seek(0)
        sheet = load_workbook(buffer).active
        self.assertEqual(sheet.cell(2, 1).value, "router")
        self.assertEqual(sheet.max_row, 2)


if __name__ == "__main__":
    unittest.main()
