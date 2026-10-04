"""Interface traffic, utilization and growth from retained Zabbix hourly trends."""
import re
from collections import defaultdict
from datetime import datetime, timedelta

RATE_UNITS = {"bps": 1, "b/s": 1, "Bps": 8, "B/s": 8}
SPEED_UNITS = {"bps": 1, "b/s": 1, "Mbps": 1e6}


def validate_config(config):
    settings = config.get("link_usage", {}) or {}
    pattern = settings.get("description_regex", "PRI|SEC")
    if not isinstance(pattern, str) or not pattern.strip():
        raise ValueError("link_usage.description_regex must be a non-empty regex")
    months = settings.get("history_months", 3)
    if type(months) is not int or not 1 <= months <= 12:
        raise ValueError("link_usage.history_months must be an integer from 1 to 12")
    coverage = settings.get("growth_min_coverage_pct", 90)
    if not isinstance(coverage, (int, float)) or not 0 <= coverage <= 100:
        raise ValueError("link_usage.growth_min_coverage_pct must be between 0 and 100")
    try:
        compiled = re.compile(pattern, re.IGNORECASE)
    except re.error as exc:
        raise ValueError(f"Invalid link_usage.description_regex: {exc}") from exc
    return settings, compiled


def calendar_months(end, count):
    """Complete calendar months ending with the month containing end (inclusive)."""
    month = datetime.fromtimestamp(end).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    periods = []
    for _ in range(count):
        next_month = (month.replace(day=28) + timedelta(days=4)).replace(day=1)
        periods.append((int(month.timestamp()), int(next_month.timestamp())))
        month = (month - timedelta(days=1)).replace(day=1)
    return list(reversed(periods))


def discover_interfaces(api, host_id, pattern):
    items = api._call("item.get", {
        "hostids": [host_id], "output": ["itemid", "key_", "name", "units", "value_type", "lastvalue"],
        "filter": {"status": 0},
    })
    interfaces = {}
    for item in items:
        match = re.fullmatch(r'net\.if\.(in|out|speed|alias|descr)\[(.+)\]', item["key_"])
        if not match:
            continue
        kind, parameter = match.groups()
        parameter = parameter.split(",")[0].strip('"')
        identity = parameter.rsplit(".", 1)[-1] if re.match(r'^if[A-Za-z]+\.', parameter) else parameter
        interface = interfaces.setdefault(identity, {})
        interface[kind] = item
        name = re.fullmatch(r'Interface (.+?)\((.*)\): .+', item["name"])
        if name:
            interface["name"], interface["description"] = name.groups()
    selected = {}
    for identity, interface in interfaces.items():
        alias = interface.get("alias") or interface.get("descr") or {}
        description = alias.get("lastvalue") or interface.get("description", "")
        if description and pattern.search(description):
            interface["description"] = description
            selected[identity] = interface
    return selected


def fetch_trends(api, interfaces, start, stop):
    ids = list(dict.fromkeys(item["itemid"] for interface in interfaces.values()
               for kind in ("in", "out", "speed") if (item := interface.get(kind))
               and item.get("units") in (SPEED_UNITS if kind == "speed" else RATE_UNITS)))
    result = defaultdict(list)
    # Keep individual responses bounded while reducing request overhead.
    for offset in range(0, len(ids), 12):
        batch = ids[offset:offset + 12]
        trends = api._call("trend.get", {"itemids": batch, "time_from": start, "time_till": stop - 1,
            "output": ["itemid", "clock", "num", "value_avg", "value_max"]})
        for trend in trends:
            if int(trend["num"]) > 0:
                result[str(trend["itemid"])].append(trend)
    return result


def summarize(trends, start, stop, factor):
    bits, covered, peak = 0.0, 0, None
    for trend in trends:
        seconds = max(0, min(int(trend["clock"]) + 3600, stop) - max(int(trend["clock"]), start))
        if not seconds:
            continue
        covered += seconds
        bits += float(trend["value_avg"]) * factor * seconds
        peak = max(peak or 0, float(trend["value_max"]) * factor)
    return {"average": bits / covered if covered else None, "peak": peak,
            "gb": bits / 8 / 1e9 if covered else None, "coverage": covered / (stop - start) * 100}


def percent_change(previous, current):
    return (current - previous) / previous * 100 if previous is not None and previous > 0 and current is not None else None


def build_rows(interfaces, trends, periods, daily=False, min_coverage=90):
    rows = []
    for identity, interface in interfaces.items():
        speed = interface.get("speed", {})
        speed_factor = SPEED_UNITS.get(speed.get("units"), 0)
        current_capacity = float(speed.get("lastvalue") or 0) * speed_factor
        # Partition once; avoid scanning three months of hourly data for each day.
        buckets = {}
        for kind in ("in", "out", "speed"):
            item = interface.get(kind, {})
            grouped = defaultdict(list)
            for trend in trends.get(str(item.get("itemid")), []):
                date = datetime.fromtimestamp(int(trend["clock"]))
                grouped[date.strftime("%Y-%m-%d" if daily else "%Y-%m")].append(trend)
            buckets[kind] = grouped
        previous = None
        for start, stop in periods:
            key = datetime.fromtimestamp(start).strftime("%Y-%m-%d" if daily else "%Y-%m")
            historical = summarize(buckets["speed"].get(key, []), start, stop, speed_factor)
            historical_capacity = historical["average"]
            capacity = historical_capacity if historical_capacity and historical_capacity > 0 else current_capacity
            row = {"Interface": interface.get("name", identity), "Interface ID": identity,
                   "Description": interface["description"], "Date" if daily else "Month": key,
                   "Capacity (Mbps)": capacity / 1e6 if capacity > 0 else None,
                   "Historical Capacity (Mbps)": historical_capacity / 1e6 if historical_capacity and historical_capacity > 0 else None,
                   "Capacity Basis": "Historical interface speed" if historical_capacity and historical_capacity > 0 else "Current interface speed",
                   "Capacity Trend Coverage (%)": historical["coverage"]}
            for kind, label in (("in", "Download"), ("out", "Upload")):
                factor = RATE_UNITS.get(interface.get(kind, {}).get("units"), 0)
                stats = summarize(buckets[kind].get(key, []), start, stop, factor)
                average, peak = stats["average"], stats["peak"]
                row.update({f"{label} (GB)": stats["gb"],
                    f"{label} (TB)": stats["gb"] / 1000 if stats["gb"] is not None else None,
                    f"{label} Avg (Mbps)": average / 1e6 if average is not None else None,
                    f"{label} Peak (Mbps)": peak / 1e6 if peak is not None else None,
                    f"{label} Avg (%)": average / capacity * 100 if average is not None and capacity > 0 else None,
                    f"{label} Peak (%)": peak / capacity * 100 if peak is not None and capacity > 0 else None,
                    f"{label} Trend Coverage (%)": stats["coverage"]})
                if not daily:
                    comparable = previous is not None and stats["coverage"] >= min_coverage and previous[f"{label} Trend Coverage (%)"] >= min_coverage
                    for metric in ("Avg", "Peak"):
                        row[f"{label} {metric} Growth (%)"] = percent_change(previous[f"{label} {metric} (Mbps)"], row[f"{label} {metric} (Mbps)"]) if comparable else None
            if not daily:
                comparable_capacity = previous is not None and historical["coverage"] >= min_coverage and previous["Capacity Trend Coverage (%)"] >= min_coverage and row["Capacity Basis"] == previous["Capacity Basis"] == "Historical interface speed"
                row["Capacity Growth (%)"] = percent_change(previous["Capacity (Mbps)"], row["Capacity (Mbps)"]) if comparable_capacity else None
            rows.append(row)
            previous = row
    return rows


def collect_link_report(api, host_id, start, end, config):
    settings, pattern = validate_config(config)
    if not settings.get("enabled", True):
        return {"latest": [], "monthly": [], "daily": []}
    periods = calendar_months(end, settings.get("history_months", 3))
    interfaces = discover_interfaces(api, host_id, pattern)
    trends = fetch_trends(api, interfaces, periods[0][0], periods[-1][1])
    monthly = build_rows(interfaces, trends, periods, min_coverage=settings.get("growth_min_coverage_pct", 90))
    days = []
    day = datetime.fromtimestamp(periods[0][0])
    while day.timestamp() < periods[-1][1]:
        following = day + timedelta(days=1)
        days.append((int(day.timestamp()), int(following.timestamp())))
        day = following
    daily = build_rows(interfaces, trends, days, daily=True)
    latest_month = datetime.fromtimestamp(periods[-1][0]).strftime("%Y-%m")
    return {"latest": [row for row in monthly if row["Month"] == latest_month], "monthly": monthly, "daily": daily}


def monthly_link_usage(api, host_id, start, end, config):
    """Single-range compatibility helper."""
    settings, pattern = validate_config(config)
    if not settings.get("enabled", True):
        return []
    interfaces = discover_interfaces(api, host_id, pattern)
    trends = fetch_trends(api, interfaces, start, end + 1)
    return build_rows(interfaces, trends, [(start, end + 1)])


def usage_rows(hosts, field="link_usage"):
    return [{"Host": host["name"], **row} for host in hosts for row in host.get(field, [])]


def add_usage_sheet(report, group_name, hosts):
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.series import SeriesLabel
    for prefix, field in (("Links", "link_usage"), ("Monthly", "link_history"), ("Daily", "link_daily")):
        rows = usage_rows(hosts, field)
        if not rows:
            continue
        title = re.sub(r'[\\/*?:\[\]]', '-', f"{prefix} {group_name}")[:31]
        sheet = report.workbook.create_sheet(title)
        sheet.append(list(rows[0]))
        for row in rows:
            sheet.append(list(row.values()))
        for cell in sheet[1]:
            cell.font = report.header_font
            cell.fill = report.header_fill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(60, max(20, len(str(column[0].value)) + 2))
            for cell in column[1:]:
                if isinstance(cell.value, (float, int)):
                    cell.number_format = "0.00"
        if prefix == "Monthly":
            # Each interface's contiguous month rows get volume, peak, and capacity charts.
            headers = list(rows[0])
            blocks = defaultdict(list)
            for index, row in enumerate(rows, 2):
                blocks[(row["Host"], row["Interface ID"])].append(index)
            for block_number, ((host, identity), indices) in enumerate(blocks.items()):
                for chart_number, (label, columns) in enumerate((
                    ("Traffic (GB)", ["Download (GB)", "Upload (GB)"]),
                    ("Peak usage (Mbps)", ["Download Peak (Mbps)", "Upload Peak (Mbps)"]),
                    ("Capacity and demand (Mbps)", ["Historical Capacity (Mbps)", "Download Avg (Mbps)", "Upload Avg (Mbps)"]),
                )):
                    chart = LineChart()
                    chart.title = f"{host} / {identity}: {label}"
                    chart.y_axis.title = label
                    chart.display_blanks = "gap"
                    for column in columns:
                        chart.add_data(Reference(sheet, min_col=headers.index(column) + 1, min_row=min(indices), max_row=max(indices)))
                        chart.series[-1].title = SeriesLabel(v=column)
                    chart.set_categories(Reference(sheet, min_col=headers.index("Month") + 1, min_row=min(indices), max_row=max(indices)))
                    sheet.add_chart(chart, f"{get_column_letter(len(headers) + 2 + chart_number * 10)}{1 + block_number * 16}")
