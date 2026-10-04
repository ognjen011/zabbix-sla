"""Retain standalone monthly exports in the app's existing report history."""
from pathlib import Path

from openpyxl import load_workbook

import database as db


def sheet_rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    return [dict(zip(headers, row)) for row in iterator]


def retain_monthly_workbook(path, group_name, generated_by='CLI'):
    path = Path(path)
    db.init_db()
    existing_id = next((existing["id"] for existing in db.get_reports(limit=db.get_report_count()) if existing["report_name"] == path.name), None)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sla = sheet_rows(next(sheet for sheet in workbook if sheet.title.startswith('SLA ') and sheet.title != 'SLA Trends'))
        sla = [row for row in sla if row.get('Technical Host')]
        latest = max(row['Month'] for row in sla)
        hosts = {row['Host']: {'name':row['Host'], 'host':row['Technical Host'],
                 'device_sla':row['Availability (%)'], 'avail_prev_month':row['Availability (%)'],
                 'sla_status':row['Status'], 'link_usage':[], 'link_history':[], 'link_daily':[]}
                 for row in sla if row['Month'] == latest}
        for host_name, host in hosts.items():
            host['sla_history'] = [{'Month':row['Month'], 'Availability (%)':row['Availability (%)'],
                'SLA Target (%)':row['SLA Target (%)'], 'Status':row['Status'],
                'Downtime (seconds)':row.get('Downtime (seconds)')}
                for row in sla if row['Host'] == host_name]
        has_monthly = any(sheet.title.startswith('Monthly ') for sheet in workbook)
        for sheet in workbook:
            if sheet.title.startswith('Links '):
                fields = ['link_usage'] if has_monthly else ['link_usage', 'link_history']
            elif sheet.title.startswith('Monthly '):
                fields = ['link_history']
            elif sheet.title.startswith('Daily '):
                fields = ['link_daily']
            else:
                continue
            for row in sheet_rows(sheet):
                host = hosts.get(row.pop('Host', None))
                if host is not None:
                    for field in fields:
                        host[field].append(row)
        latest_rows = [row for row in sla if row['Month'] == latest]
        months = sorted({row['Month'] for row in sla})
        summary = {'group_name':group_name, 'sla_threshold':latest_rows[0]['SLA Target (%)'],
                   'total':len(hosts), **{status.lower():sum(row['Status'] == status for row in latest_rows) for status in ('COMPLIANT','WARNING','BREACH')}}
        if 'SLA Trends' in workbook.sheetnames:
            group_history = [{'Month':row['Month'], 'Availability (%)':row['Overall SLA (%)'],
                'SLA Target (%)':row.get('SLA Target (%)', summary['sla_threshold']),
                'Compliant':row.get('Compliant'), 'Warning':row.get('Warning'), 'Breach':row.get('Breach')}
                for row in sheet_rows(workbook['SLA Trends'])]
            overall = next(row['Availability (%)'] for row in group_history if row['Month'] == latest)
        else:
            overall = next((row['Availability (%)'] for row in sheet_rows(next(sheet for sheet in workbook if sheet.title.startswith('SLA ') and sheet.title != 'SLA Trends')) if row['Host'] == 'Overall group'), None)
        if 'SLA Trends' not in workbook.sheetnames:
            group_history = [{'Month':latest, 'Availability (%)':overall, 'SLA Target (%)':summary['sla_threshold']}]
        summary.update({'overall_sla':overall, 'overall_prev_month':overall, 'sla_history':group_history})
        if existing_id is not None:
            db.update_report_data(existing_id, [summary], {group_name:list(hosts.values())})
            return existing_id
        return db.save_report(generated_by, path.name, 'month' if len(months) == 1 else f'{len(months)} months',
                              [group_name], len(hosts), [summary], {group_name:list(hosts.values())}, path.read_bytes())
    finally:
        workbook.close()
